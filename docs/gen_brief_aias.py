#!/usr/bin/env python3
"""Genera il brief aggiornato per l'incontro con il direttore AIAS."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Style base
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# Colore accento
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = ACCENT
    r.font.size = Pt(18)
    return p

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = ACCENT
    r.font.size = Pt(13)
    return p

def body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

# Copertina
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Brief per l'incontro con AIAS")
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = ACCENT

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st.add_run("Osservatorio Infortuni: come è costruito, cosa mostriamo, dove può crescere")
r.font.size = Pt(13)
r.font.italic = True

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Documento di supporto per il resoconto del progetto\nAggiornato al 15 ottobre 2026 · Ing. Damiano Salvati")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_page_break()

# 1. Contesto
h1("1. Contesto dell'incontro")
body("L'Osservatorio Infortuni nasce come progetto civico e professionale per rendere leggibili i dati sugli infortuni sul lavoro in Italia. La banca dati INAIL è ricchissima ma di difficile lettura: dataset enormi, codici al posto dei nomi, nessuna vista sintetica. Il progetto la valorizza con grafici e analisi che permettono una lettura facile e approfondita, sempre con fonte e data dichiarate.")
body("Il progetto è aperto, gratuito e indipendente. I dati sono pubblici (INAIL, INL, Eurostat, ISTAT), ogni numero pubblicato è verificato e riproducibile. L'obiettivo dell'incontro è raccontare come è fatto, mostrare cosa funziona già e capire se e come può servire i programmi AIAS/ENSHPO 2026-2030.")

# 2. Come è costruito
h1("2. Come è costruito il progetto")
body("Architettura tecnica: sito web open source (Next.js) con pipeline di dati in Python. I dati non vengono interrogati dal browser in tempo reale: il sito pubblica snapshot verificati, rigenerati con procedure ETL controllate. Un audit automatico controlla gli endpoint delle fonti ufficiali il 15 di ogni mese e segnala i cambiamenti, senza introdurre numeri senza revisione.")
bullet("Base dati INAIL: serie storica consolidata 2021-2025, dataset semestrali regionali, tutte le 20 regioni, circa 3,05 milioni di record (rilevazione 30/04/2026).")
bullet("Congiuntura mensile 2025-2026: API REST mensile e CSV, circa 493.000 casi nella finestra mobile.")
bullet("Analisi Infor.MO (INAIL): archivio delle cause, con dinamiche e fattori causali aggregati, incrocio causa per settore economico.")
bullet("Malattie professionali: denunce e decessi riconosciuti, con perimetro dichiarato nella pagina.")
bullet("Vigilanza: rapporti annuali INL e indicatori sull'attività ispettiva.")
bullet("Radar degli infortuni: rassegna quotidiana da fonti di cronaca, filtrata con un punteggio interno di pertinenza. Non è una statistica ufficiale, è una cronaca per non dimenticare.")
body("Principio guida: nessun numero senza fonte e data. Ogni pagina dichiara il proprio perimetro, la data del dato e le discontinuità note. Un aumento delle denunce può dipendere da più vigilanza, più consapevolezza o cambi normativi: i confronti usano solo misure compatibili.")

# 3. Cosa mostriamo oggi
h1("3. Cosa mostriamo oggi (numeri verificati al 15/10)")
h2("Infortuni sul lavoro")
bullet("3.052.546 denunce nel quinquennio 2021-2025, 20 regioni, 107 province.")
bullet("6.429 casi mortali nel periodo (dato da rilevazione INAIL, estrazione 23/09/2026).")
bullet("Serie storiche per tempo, luogo (regione e provincia), persona (età, genere), modalità (in occasione di lavoro e in itinere), settore (ATECO) ed esito.")
bullet("Benchmark europeo Eurostat ESAW per contestualizzare l'Italia.")
h2("Radar di cronaca (nuovo, sezione rinnovata a ottobre)")
bullet("Archivio progressivo fino a 365 giorni, oggi 229 notizie raccolte e classificate.")
bullet("Tre categorie: 141 mortali, 43 gravi, 45 analisi e approfondimenti.")
bullet("Racconta i fatti con sintesi brevi, tutte le testate per ogni evento, ricerca per parole chiave. Niente numeri KPI in questa sezione: conta il racconto, non lo score.")
h2("Malattie professionali (sezione aggiornata)")
bullet("110.279 denunce nel confronto a pari periodo: primo semestre 2026 a 59.394 casi contro 50.885 nel 2025, +16,7% anno su anno.")
bullet("4.853 decessi riconosciuti nell'intero dataset, con dettaglio per patologia.")
bullet("Esiti 2025: 98.312 denunce, 36,6% riconosciute, 50,7% rigettate, 12,7% in istruttoria.")
h2("Vigilanza")
bullet("Serie storica annuale dell'attività INL (2021-2025): controlli avviati, ispezioni definite, tasso di irregolarità, patente a crediti.")

# 4. LAB Sicurezza Industriale
h1("4. Il laboratorio Sicurezza Industriale (in sviluppo)")
body("A fianco dell'Osservatorio è in costruzione un laboratorio per i professionisti della sicurezza. L'idea: unire ai dati istituzionali gli strumenti di analisi del rischio che oggi richiedono software costosi.")
bullet("Atlante delle banche dati di affidabilità (IOGP, HSE FRED, OREDA, CCPS, HEART e altre, 14 fonti classificate tra dati aperti, link e a pagamento).")
bullet("Database consultabile con 25 voci di frequenze di guasto e probabilità, 6 categorie (serbatoi, tubazioni, macchine rotanti, valvole, scambiatori, fattore umano), ogni valore con fonte e anno.")
bullet("Lavagna FTA/ETA in stile Miro: si costruiscono alberi di guasto e alberi di eventi con porte logiche, si assegnano le probabilità dal database con un click e si ottengono cut set e probabilità in tempo reale.")
bullet("Il laboratorio è pubblicato su ambiente separato e sperimentale. Serve a dimostrare che il flusso professionista completo, dalla banca dati al calcolo, è già percorribile.")

# 5. Cosa proponiamo ad AIAS
h1("5. Cosa possiamo offrire ad AIAS/ENSHPO")
body("Tre livelli di collaborazione, tutti reversibili e compatibili con l'indipendenza del progetto:")
bullet("Dati con fonte: condividere le basi verificate dell'Osservatorio per i programmi 2026-2030, sempre con metodologia e limiti dichiarati.")
bullet("Co-branding di report: analisi congiunte su temi del programma (salute mentale, PMI, qualifiche OSH), con il marchio dell'Osservatorio e di AIAS.")
bullet("Advisory sul metodo: il progetto come banco di prova per come leggere i dati, costruire indicatori e comunicare i risultati senza fuorviare.")
body("In cambio chiediamo visibilità, eventuale patrocinio e un ruolo riconosciuto nel programma. Il progetto resta pubblico e indipendente: i dati sono di tutti, il valore sta nella catena che li rende leggibili.")

# 6. Cosa chiediamo e cosa ascoltiamo
h1("6. Domande che portiamo all'incontro")
body("Prima di tutto ascoltiamo: quali dati servono ai programmi 2026-2030, per cosa e per quando. Poi chiarimenti operativi:")
bullet("Quali temi del programma 2027-2030 hanno bisogno di dati e analisi?")
bullet("Come si inquadra una collaborazione con un progetto civico open source?")
bullet("C'è spazio per una sperimentazione concreta entro l'assemblea di novembre o nei primi mesi del 2027?")
bullet("Quali vincoli di immagine e di metodo per l'uso dei dati da parte dell'associazione?")
body("Linea in una frase: raccontiamo un'infrastruttura dati già pronta, mostriamo cosa funziona, e proponiamo una collaborazione a tre livelli che rispetta l'indipendenza del progetto e dà valore ai programmi AIAS/ENSHPO.")

# 7. Note operative
h1("7. Note operative")
bullet("Sito pubblico: osservatorioinfortuni.it (Vercel, deploy automatico da GitHub).")
bullet("Licenza AGPL-3.0, codice sorgente aperto su GitHub.")
bullet("Laboratorio Sicurezza Industriale: ambiente separato, in sviluppo.")
bullet("Regola privacy: nessun caso singolo pubblicato, solo aggregazioni.")
bullet("Dopo l'incontro: inviare questo documento come promemoria e annotare i prossimi passi.")

doc.save("/root/.openclaw/workspace/projects/osservatorio-infortuni/docs/Brief-incontro-AIAS-2026-10-15.docx")
print("Salvato OK")
